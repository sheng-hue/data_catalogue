"""
Snowflake Data Catalog Documentation Generator
A Streamlit app for generating LLM-powered column documentation for Snowflake views.
Handles 200-1200 columns efficiently with batching and one-pass sampling.
"""

import streamlit as st
import pandas as pd
import snowflake.connector
import json
import re
import time
import os
from typing import Dict, List, Tuple, Optional, Any
from datetime import datetime, date
from concurrent.futures import ThreadPoolExecutor, as_completed
from io import BytesIO
import traceback

# Configure page
st.set_page_config(
    page_title="Snowflake Data Catalog",
    page_icon="❄️",
    layout="wide",
    initial_sidebar_state="expanded"
)

# Custom CSS for better UX
st.markdown("""
<style>
    .stProgress > div > div > div > div {
        background-color: #1f77b4;
    }
    .success-box {
        padding: 1rem;
        border-radius: 0.5rem;
        background-color: #d4edda;
        border: 1px solid #c3e6cb;
    }
    .warning-box {
        padding: 1rem;
        border-radius: 0.5rem;
        background-color: #fff3cd;
        border: 1px solid #ffeaa7;
    }
    .error-box {
        padding: 1rem;
        border-radius: 0.5rem;
        background-color: #f8d7da;
        border: 1px solid #f5c6cb;
    }
</style>
""", unsafe_allow_html=True)

# ============================================================================
# SESSION STATE INITIALIZATION
# ============================================================================

def init_session_state():
    """Initialize all session state variables"""
    defaults = {
        'connected': False,
        'connection': None,
        'databases': [],
        'schemas': [],
        'views': [],
        'selected_db': None,
        'selected_schema': None,
        'selected_view': None,
        'columns_data': None,
        'sample_data': None,
        'results_df': None,
        'generation_complete': False,
        'failed_batches': [],
        'llm_cache': {}  # Cache for LLM results if opted in
    }

    for key, value in defaults.items():
        if key not in st.session_state:
            st.session_state[key] = value

init_session_state()

# ============================================================================
# SNOWFLAKE CONNECTION FUNCTIONS
# ============================================================================

def parse_account(input_str: str) -> str:
    """
    Parse Snowflake account from URL or account identifier.

    Args:
        input_str: Either a URL (https://...snowflakecomputing.com) or account identifier

    Returns:
        Account identifier string
    """
    input_str = input_str.strip()

    # If it's a URL, extract account from it
    if input_str.startswith('http://') or input_str.startswith('https://'):
        # Extract domain
        pattern = r'https?://([^/]+)'
        match = re.search(pattern, input_str)
        if match:
            domain = match.group(1)
            # Remove .snowflakecomputing.com
            account = domain.replace('.snowflakecomputing.com', '')
            return account

    # Otherwise assume it's already an account identifier
    return input_str

@st.cache_resource
def get_snowflake_connection(account: str, user: str, role: Optional[str] = None,
                             warehouse: Optional[str] = None) -> snowflake.connector.SnowflakeConnection:
    """
    Create and cache a Snowflake connection using SSO (external browser authentication).

    Args:
        account: Snowflake account identifier
        user: User email/username
        role: Optional role to use
        warehouse: Optional warehouse to use

    Returns:
        Snowflake connection object
    """
    connection_params = {
        'account': account,
        'user': user,
        'authenticator': 'externalbrowser',
    }

    if role:
        connection_params['role'] = role
    if warehouse:
        connection_params['warehouse'] = warehouse

    try:
        conn = snowflake.connector.connect(**connection_params)
        return conn
    except Exception as e:
        st.error(f"Connection failed: {str(e)}")
        raise

def get_connection_info(conn: snowflake.connector.SnowflakeConnection) -> Dict[str, str]:
    """Get current connection information"""
    cursor = conn.cursor()

    info = {}
    try:
        cursor.execute("SELECT CURRENT_USER()")
        info['user'] = cursor.fetchone()[0]

        cursor.execute("SELECT CURRENT_ROLE()")
        info['role'] = cursor.fetchone()[0]

        cursor.execute("SELECT CURRENT_WAREHOUSE()")
        result = cursor.fetchone()
        info['warehouse'] = result[0] if result else "None"
    finally:
        cursor.close()

    return info

# ============================================================================
# SNOWFLAKE METADATA FUNCTIONS
# ============================================================================

@st.cache_data(ttl=300)
def list_databases(_conn: snowflake.connector.SnowflakeConnection) -> List[str]:
    """List all available databases"""
    cursor = _conn.cursor()
    try:
        cursor.execute("SHOW DATABASES")
        databases = [row[1] for row in cursor.fetchall()]  # name is in column 1
        return sorted(databases)
    finally:
        cursor.close()

@st.cache_data(ttl=300)
def list_schemas(_conn: snowflake.connector.SnowflakeConnection, database: str) -> List[str]:
    """List all schemas in a database"""
    cursor = _conn.cursor()
    try:
        cursor.execute(f"SHOW SCHEMAS IN DATABASE {database}")
        schemas = [row[1] for row in cursor.fetchall()]  # name is in column 1
        return sorted(schemas)
    finally:
        cursor.close()

@st.cache_data(ttl=300)
def list_views(_conn: snowflake.connector.SnowflakeConnection, database: str, schema: str) -> List[str]:
    """List all views in a schema"""
    cursor = _conn.cursor()
    try:
        cursor.execute(f"SHOW VIEWS IN {database}.{schema}")
        views = [row[1] for row in cursor.fetchall()]  # name is in column 1
        return sorted(views)
    finally:
        cursor.close()

@st.cache_data(ttl=300)
def get_view_columns(_conn: snowflake.connector.SnowflakeConnection,
                     database: str, schema: str, view: str) -> pd.DataFrame:
    """
    Get column information for a view using INFORMATION_SCHEMA.

    Returns DataFrame with: column_name, data_type, ordinal_position
    """
    query = f"""
    SELECT
        COLUMN_NAME,
        DATA_TYPE,
        ORDINAL_POSITION
    FROM {database}.INFORMATION_SCHEMA.COLUMNS
    WHERE TABLE_SCHEMA = '{schema}'
      AND TABLE_NAME = '{view}'
    ORDER BY ORDINAL_POSITION
    """

    cursor = _conn.cursor()
    try:
        cursor.execute(query)
        columns_data = cursor.fetchall()
        df = pd.DataFrame(columns_data, columns=['column_name', 'data_type', 'ordinal_position'])
        return df
    finally:
        cursor.close()

# ============================================================================
# SAMPLING FUNCTIONS
# ============================================================================

def format_sample_value(value: Any, data_type: str, max_length: int = 120) -> str:
    """Format a sample value based on its data type"""
    if value is None or pd.isna(value):
        return ""

    try:
        # Handle VARIANT, OBJECT, ARRAY types
        if data_type in ['VARIANT', 'OBJECT', 'ARRAY']:
            if isinstance(value, (dict, list)):
                json_str = json.dumps(value)
            else:
                json_str = str(value)

            if len(json_str) > max_length:
                return json_str[:max_length] + "..."
            return json_str

        # Handle dates and timestamps
        if isinstance(value, (datetime, date)):
            return value.isoformat()

        # Handle strings
        if isinstance(value, str):
            if len(value) > max_length:
                return value[:max_length] + "..."
            return value

        # Handle numbers and other types
        return str(value)
    except Exception as e:
        return f"<error: {str(e)[:50]}>"

@st.cache_data(ttl=600)
def sample_rows_onepass(_conn: snowflake.connector.SnowflakeConnection,
                        database: str, schema: str, view: str,
                        n_rows: int = 2000) -> pd.DataFrame:
    """
    Perform one-pass sampling of the view.
    This is much more efficient than per-column sampling for wide tables.

    Args:
        _conn: Snowflake connection
        database, schema, view: View identifier
        n_rows: Number of rows to sample

    Returns:
        DataFrame with sampled rows
    """
    query = f"""
    SELECT * FROM {database}.{schema}.{view}
    SAMPLE ({n_rows} ROWS)
    LIMIT {n_rows}
    """

    cursor = _conn.cursor()
    try:
        cursor.execute(query)
        # Fetch all rows
        rows = cursor.fetchall()
        column_names = [desc[0] for desc in cursor.description]

        # Create DataFrame
        df = pd.DataFrame(rows, columns=column_names)
        return df
    finally:
        cursor.close()

def collect_samples_from_df(df: pd.DataFrame, columns_meta: pd.DataFrame,
                            max_values: int = 15) -> Dict[str, List[str]]:
    """
    Collect up to max_values distinct non-null samples for each column from the DataFrame.

    Args:
        df: Sampled data DataFrame
        columns_meta: DataFrame with column metadata (column_name, data_type)
        max_values: Maximum number of sample values to collect per column

    Returns:
        Dictionary mapping column_name to list of formatted sample strings
    """
    samples_map = {}

    # Create a dict for quick data type lookup
    dtype_map = dict(zip(columns_meta['column_name'], columns_meta['data_type']))

    for col in df.columns:
        data_type = dtype_map.get(col, 'TEXT')

        # Get non-null values
        non_null_values = df[col].dropna()

        if len(non_null_values) == 0:
            samples_map[col] = []
            continue

        # Try to get distinct values (but don't fail if unhashable)
        try:
            unique_values = non_null_values.unique()
        except TypeError:
            # If values are unhashable, just take first N
            unique_values = non_null_values.values

        # Take up to max_values
        sample_values = unique_values[:max_values]

        # Format each value
        formatted_samples = [
            format_sample_value(val, data_type)
            for val in sample_values
        ]

        # Remove empty strings
        formatted_samples = [s for s in formatted_samples if s]

        samples_map[col] = formatted_samples[:max_values]

    return samples_map

def sample_rows_chunked(_conn: snowflake.connector.SnowflakeConnection,
                        database: str, schema: str, view: str,
                        columns: List[str], data_types: Dict[str, str],
                        max_values: int = 15, max_workers: int = 8) -> Dict[str, List[str]]:
    """
    Fallback: sample each column individually with parallel execution.
    This is slower but works when one-pass sampling fails.

    Args:
        _conn: Snowflake connection
        database, schema, view: View identifier
        columns: List of column names
        data_types: Dict mapping column name to data type
        max_values: Number of samples per column
        max_workers: Max parallel workers

    Returns:
        Dictionary mapping column_name to list of sample strings
    """
    def sample_column(col_name: str) -> Tuple[str, List[str]]:
        """Sample a single column"""
        query = f"""
        SELECT DISTINCT "{col_name}"
        FROM {database}.{schema}.{view}
        WHERE "{col_name}" IS NOT NULL
        LIMIT {max_values * 2}
        """

        cursor = _conn.cursor()
        try:
            cursor.execute(query)
            values = [row[0] for row in cursor.fetchall()]

            data_type = data_types.get(col_name, 'TEXT')
            formatted = [format_sample_value(v, data_type) for v in values]
            formatted = [s for s in formatted if s]

            return col_name, formatted[:max_values]
        except Exception as e:
            return col_name, [f"<error: {str(e)[:50]}>"]
        finally:
            cursor.close()

    samples_map = {}

    # Use ThreadPoolExecutor for parallel sampling
    with ThreadPoolExecutor(max_workers=max_workers) as executor:
        futures = {executor.submit(sample_column, col): col for col in columns}

        for future in as_completed(futures):
            col_name, samples = future.result()
            samples_map[col_name] = samples

    return samples_map

# ============================================================================
# LLM FUNCTIONS
# ============================================================================

def call_cortex_llm(conn: snowflake.connector.SnowflakeConnection,
                   model: str, prompt: str, max_retries: int = 2) -> str:
    """
    Call Snowflake Cortex LLM function using parameterized query.

    Args:
        conn: Snowflake connection
        model: Model name (e.g., 'mistral-large', 'mixtral-8x7b', 'llama2-70b-chat')
        prompt: Prompt text
        max_retries: Number of retries on failure

    Returns:
        LLM response text
    """
    # Use parameterized query to avoid SQL injection and escaping issues
    query = """
    SELECT SNOWFLAKE.CORTEX.COMPLETE(%s, %s) AS response
    """

    cursor = conn.cursor()
    try:
        for attempt in range(max_retries + 1):
            try:
                cursor.execute(query, (model, prompt))
                result = cursor.fetchone()
                return result[0] if result else ""
            except Exception as e:
                if attempt == max_retries:
                    raise
                time.sleep(2 ** attempt)  # Exponential backoff
    finally:
        cursor.close()

def build_llm_prompt(columns_batch: pd.DataFrame,
                    view_description: str,
                    samples_map: Dict[str, List[str]],
                    settings: Dict[str, Any]) -> str:
    """
    Build the LLM prompt for a batch of columns.

    Args:
        columns_batch: DataFrame with column_name, data_type
        view_description: User-provided view description
        samples_map: Dict of column -> sample values
        settings: Generation settings dict

    Returns:
        Formatted prompt string
    """
    style = settings.get('output_style', 'technical')
    max_length = settings.get('max_length', 200)
    include_types = settings.get('include_types', True)
    include_samples = settings.get('include_samples', True)
    include_view_desc = settings.get('include_view_desc', True)

    # Style instructions
    style_instructions = {
        'business': "Write in business-friendly, non-technical language. Avoid jargon. Focus on what the data means for business users.",
        'technical': "Write in technical, data-engineering language. Include relevant technical details and data engineering context."
    }

    style_inst = style_instructions.get(style, style_instructions['technical'])

    prompt_parts = [
        "You are a data documentation expert. Generate concise, accurate descriptions for the following database columns.",
        f"\n{style_inst}",
        f"\nEach description must be maximum {max_length} characters.",
        "\nReturn ONLY valid JSON array with this exact format:",
        '[{"column_name":"COL1", "description":"...", "confidence":1-5}, ...]',
        "\nDo not include any text before or after the JSON array.\n"
    ]

    # Add view context if enabled
    if include_view_desc and view_description:
        prompt_parts.append(f"\nVIEW CONTEXT:\n{view_description}\n")

    # Add column information
    prompt_parts.append("\nCOLUMNS TO DOCUMENT:\n")

    for idx, row in columns_batch.iterrows():
        col_name = row['column_name']
        data_type = row['data_type']

        col_info = f"\n{col_name}"

        if include_types:
            col_info += f" ({data_type})"

        if include_samples and col_name in samples_map:
            samples = samples_map[col_name]
            if samples:
                samples_str = ", ".join(str(s) for s in samples[:5])  # Limit samples in prompt
                col_info += f"\n  Sample values: {samples_str}"

        prompt_parts.append(col_info)

    prompt_parts.append("\n\nReturn the JSON array now:")

    return "".join(prompt_parts)

def parse_llm_response(response: str) -> List[Dict[str, Any]]:
    """
    Parse LLM JSON response and validate structure.

    Args:
        response: Raw LLM response text

    Returns:
        List of dicts with column_name, description, confidence

    Raises:
        ValueError if parsing fails
    """
    # Try to extract JSON from response
    response = response.strip()

    # Find JSON array in response
    start_idx = response.find('[')
    end_idx = response.rfind(']')

    if start_idx == -1 or end_idx == -1:
        raise ValueError("No JSON array found in response")

    json_str = response[start_idx:end_idx + 1]

    try:
        parsed = json.loads(json_str)
    except json.JSONDecodeError as e:
        raise ValueError(f"Invalid JSON: {str(e)}")

    # Validate structure
    if not isinstance(parsed, list):
        raise ValueError("Response is not a JSON array")

    for item in parsed:
        if not isinstance(item, dict):
            raise ValueError("Array items must be objects")
        if 'column_name' not in item or 'description' not in item:
            raise ValueError("Missing required fields: column_name, description")

    return parsed

def generate_descriptions_batch(columns_batch: pd.DataFrame,
                                view_description: str,
                                samples_map: Dict[str, List[str]],
                                settings: Dict[str, Any],
                                conn: snowflake.connector.SnowflakeConnection) -> Tuple[List[Dict], Optional[str]]:
    """
    Generate descriptions for a batch of columns using Snowflake Cortex.

    Args:
        columns_batch: DataFrame with columns to document
        view_description: View description
        samples_map: Sample values map
        settings: Generation settings
        conn: Snowflake connection (required for Cortex)

    Returns:
        Tuple of (results_list, error_message)
    """
    model = settings.get('model', 'mixtral-8x7b')

    # Build prompt
    prompt = build_llm_prompt(columns_batch, view_description, samples_map, settings)

    # Call Snowflake Cortex LLM
    try:
        response = call_cortex_llm(conn, model, prompt, max_retries=2)

        # Parse response
        try:
            results = parse_llm_response(response)
        except ValueError as parse_error:
            # Retry once with stricter prompt
            strict_prompt = prompt + "\n\nIMPORTANT: Return ONLY the JSON array. No explanations, no markdown, no code blocks. Just the JSON array starting with [ and ending with ]."
            response = call_cortex_llm(conn, model, strict_prompt, max_retries=1)
            results = parse_llm_response(response)

        # Ensure all columns are covered
        result_cols = {r['column_name'] for r in results}
        expected_cols = set(columns_batch['column_name'])

        missing_cols = expected_cols - result_cols
        if missing_cols:
            # Add placeholder for missing columns
            for col in missing_cols:
                results.append({
                    'column_name': col,
                    'description': '<LLM did not return description>',
                    'confidence': 1
                })

        return results, None

    except Exception as e:
        error_msg = f"Cortex LLM error: {str(e)}"
        # Return partial results with error markers
        results = []
        for _, row in columns_batch.iterrows():
            results.append({
                'column_name': row['column_name'],
                'description': f'<ERROR: {str(e)[:100]}>',
                'confidence': 0
            })
        return results, error_msg

def generate_all_descriptions(columns_df: pd.DataFrame,
                             view_description: str,
                             samples_map: Dict[str, List[str]],
                             settings: Dict[str, Any],
                             conn: snowflake.connector.SnowflakeConnection,
                             progress_callback=None) -> Tuple[pd.DataFrame, List[int]]:
    """
    Generate descriptions for all columns with batching and progress tracking.

    Args:
        columns_df: DataFrame with all columns
        view_description: View description
        samples_map: Sample values map
        settings: Generation settings
        conn: Snowflake connection (required for Cortex)
        progress_callback: Function to call with progress updates

    Returns:
        Tuple of (results_dataframe, failed_batch_indices)
    """
    batch_size = settings.get('batch_size', 60)
    max_columns = settings.get('max_columns', 300)

    # Apply max columns limit
    columns_to_process = columns_df.head(max_columns)

    # Split into batches
    num_batches = (len(columns_to_process) + batch_size - 1) // batch_size
    batches = [
        columns_to_process.iloc[i * batch_size:(i + 1) * batch_size]
        for i in range(num_batches)
    ]

    all_results = []
    failed_batches = []

    start_time = time.time()

    for batch_idx, batch in enumerate(batches):
        if progress_callback:
            progress_callback(batch_idx, num_batches, start_time)

        results, error = generate_descriptions_batch(
            batch, view_description, samples_map, settings, conn
        )

        if error:
            failed_batches.append(batch_idx)
            st.warning(f"Batch {batch_idx + 1}/{num_batches} failed: {error}")

        all_results.extend(results)

    # Create results DataFrame
    results_df = pd.DataFrame(all_results)

    # Merge with original columns data to get data types
    results_df = results_df.merge(
        columns_df[['column_name', 'data_type']],
        on='column_name',
        how='left'
    )

    # Add sample values as formatted strings
    results_df['sample_values'] = results_df['column_name'].apply(
        lambda col: "; ".join(samples_map.get(col, [])[:5]) if samples_map.get(col) else ""
    )

    # Reorder columns
    results_df = results_df[['column_name', 'data_type', 'description', 'confidence', 'sample_values']]

    return results_df, failed_batches

# ============================================================================
# EXPORT FUNCTIONS
# ============================================================================

def export_to_csv(df: pd.DataFrame) -> bytes:
    """Export DataFrame to CSV bytes"""
    return df.to_csv(index=False).encode('utf-8')

def export_to_excel(df: pd.DataFrame) -> bytes:
    """
    Export DataFrame to Excel with formatting.

    Args:
        df: Results DataFrame

    Returns:
        Excel file as bytes
    """
    try:
        from openpyxl import load_workbook
        from openpyxl.styles import Font, Alignment
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise ImportError("openpyxl not installed. Run: pip install openpyxl")

    output = BytesIO()

    # Write to Excel
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        df.to_excel(writer, sheet_name='columns', index=False)

    output.seek(0)

    # Load workbook for formatting
    wb = load_workbook(output)
    ws = wb['columns']

    # Freeze top row
    ws.freeze_panes = 'A2'

    # Format header
    header_font = Font(bold=True)
    for cell in ws[1]:
        cell.font = header_font

    # Set column widths and wrap text
    column_widths = {
        'A': 30,  # column_name
        'B': 20,  # data_type
        'C': 50,  # description
        'D': 12,  # confidence
        'E': 50,  # sample_values
    }

    for col_letter, width in column_widths.items():
        ws.column_dimensions[col_letter].width = width

    # Wrap text for description and sample_values
    wrap_alignment = Alignment(wrap_text=True, vertical='top')
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        row[2].alignment = wrap_alignment  # description
        row[4].alignment = wrap_alignment  # sample_values

    # Add auto-filter
    ws.auto_filter.ref = ws.dimensions

    # Save to bytes
    output = BytesIO()
    wb.save(output)
    output.seek(0)

    return output.getvalue()

# ============================================================================
# UTILITY FUNCTIONS
# ============================================================================

def estimate_tokens(prompt: str) -> int:
    """Rough token estimate (~ 4 chars per token)"""
    return len(prompt) // 4

def format_elapsed_time(seconds: float) -> str:
    """Format elapsed time in human-readable format"""
    if seconds < 60:
        return f"{seconds:.1f}s"
    elif seconds < 3600:
        minutes = seconds / 60
        return f"{minutes:.1f}m"
    else:
        hours = seconds / 3600
        return f"{hours:.1f}h"

# ============================================================================
# STREAMLIT UI
# ============================================================================

def render_sidebar():
    """Render the sidebar with connection settings"""
    st.sidebar.title("❄️ Snowflake Connection")

    st.sidebar.markdown("### Connection Settings")

    # Connection inputs
    account_input = st.sidebar.text_input(
        "Account or URL",
        value=st.session_state.get('account_input', ''),
        help="Enter Snowflake account identifier or full URL (https://...snowflakecomputing.com)",
        key='account_input'
    )

    user_email = st.sidebar.text_input(
        "Email/Username",
        value=st.session_state.get('user_email', ''),
        key='user_email'
    )

    with st.sidebar.expander("Optional Settings"):
        role = st.text_input("Role", value="", key='role_input')
        warehouse = st.text_input("Warehouse", value="", key='warehouse_input')

    # Connect button
    if st.sidebar.button("🔌 Connect (SSO)", type="primary", use_container_width=True):
        if not account_input or not user_email:
            st.sidebar.error("Please provide account and email")
            return

        try:
            with st.spinner("Opening browser for SSO authentication..."):
                account = parse_account(account_input)
                conn = get_snowflake_connection(
                    account=account,
                    user=user_email,
                    role=role if role else None,
                    warehouse=warehouse if warehouse else None
                )

                st.session_state.connection = conn
                st.session_state.connected = True

                # Get connection info
                info = get_connection_info(conn)
                st.session_state.connection_info = info

                st.sidebar.success("✅ Connected!")
                st.rerun()

        except Exception as e:
            st.sidebar.error(f"Connection failed: {str(e)}")
            st.session_state.connected = False

    # Show connection info if connected
    if st.session_state.connected:
        st.sidebar.markdown("---")
        st.sidebar.markdown("### Connected As")
        info = st.session_state.get('connection_info', {})
        st.sidebar.markdown(f"**User:** {info.get('user', 'N/A')}")
        st.sidebar.markdown(f"**Role:** {info.get('role', 'N/A')}")
        st.sidebar.markdown(f"**Warehouse:** {info.get('warehouse', 'N/A')}")

        if st.sidebar.button("🔓 Disconnect"):
            st.session_state.connected = False
            st.session_state.connection = None
            if 'connection' in st.session_state:
                try:
                    st.session_state.connection.close()
                except:
                    pass
            # Clear cache
            st.cache_data.clear()
            st.cache_resource.clear()
            st.rerun()

def render_view_selection():
    """Render view selection interface"""
    st.header("📊 View Selection")

    conn = st.session_state.connection

    col1, col2, col3 = st.columns(3)

    with col1:
        # Database dropdown
        databases = list_databases(conn)
        selected_db = st.selectbox(
            "Database",
            options=databases,
            index=databases.index(st.session_state.selected_db) if st.session_state.selected_db in databases else 0,
            key='db_select'
        )

        if selected_db != st.session_state.selected_db:
            st.session_state.selected_db = selected_db
            st.session_state.selected_schema = None
            st.session_state.selected_view = None
            st.rerun()

    with col2:
        # Schema dropdown
        if selected_db:
            schemas = list_schemas(conn, selected_db)
            selected_schema = st.selectbox(
                "Schema",
                options=schemas,
                index=schemas.index(st.session_state.selected_schema) if st.session_state.selected_schema in schemas else 0,
                key='schema_select'
            )

            if selected_schema != st.session_state.selected_schema:
                st.session_state.selected_schema = selected_schema
                st.session_state.selected_view = None
                st.rerun()

    with col3:
        # View dropdown
        if selected_db and selected_schema:
            views = list_views(conn, selected_db, selected_schema)

            if views:
                selected_view = st.selectbox(
                    "View",
                    options=views,
                    index=views.index(st.session_state.selected_view) if st.session_state.selected_view in views else 0,
                    key='view_select'
                )

                if selected_view != st.session_state.selected_view:
                    st.session_state.selected_view = selected_view
                    st.session_state.columns_data = None
                    st.session_state.sample_data = None
            else:
                st.warning("No views found in this schema")
                selected_view = None

    # View description
    st.markdown("### View Description")
    view_description = st.text_area(
        "Describe the business meaning, grain, and any caveats",
        height=100,
        key='view_description',
        help="This context will be provided to the LLM to generate more accurate column descriptions"
    )

    # Load columns button
    if st.session_state.selected_view:
        if st.button("📥 Load Columns", type="primary"):
            try:
                with st.spinner("Loading column metadata..."):
                    columns_df = get_view_columns(
                        conn,
                        st.session_state.selected_db,
                        st.session_state.selected_schema,
                        st.session_state.selected_view
                    )

                    st.session_state.columns_data = columns_df
                    st.session_state.sample_data = None  # Reset sampling
                    st.session_state.results_df = None  # Reset results

                    st.success(f"✅ Loaded {len(columns_df)} columns")
                    st.rerun()

            except Exception as e:
                st.error(f"Failed to load columns: {str(e)}")
                st.code(traceback.format_exc())

def render_sampling_section():
    """Render sampling configuration"""
    if st.session_state.columns_data is None:
        return

    st.markdown("---")
    st.header("🎲 Sampling")

    num_columns = len(st.session_state.columns_data)
    st.info(f"View has **{num_columns}** columns")

    col1, col2 = st.columns([2, 1])

    with col1:
        skip_sampling = st.checkbox(
            "Skip sampling (faster LLM-only generation)",
            value=False,
            help="Skip sampling for faster processing. LLM will generate descriptions based only on column names and types."
        )

        if not skip_sampling:
            sample_size = st.slider(
                "Sample size (rows)",
                min_value=500,
                max_value=5000,
                value=2000,
                step=500,
                help="Number of rows to sample for collecting example values"
            )

    with col2:
        if st.button("🎲 Sample Data", type="primary", disabled=skip_sampling):
            try:
                with st.spinner("Sampling data... This may take a moment for wide tables."):
                    # Try one-pass sampling first
                    try:
                        sample_df = sample_rows_onepass(
                            st.session_state.connection,
                            st.session_state.selected_db,
                            st.session_state.selected_schema,
                            st.session_state.selected_view,
                            n_rows=sample_size
                        )

                        # Collect samples from DataFrame
                        samples_map = collect_samples_from_df(
                            sample_df,
                            st.session_state.columns_data,
                            max_values=15
                        )

                        st.session_state.sample_data = samples_map
                        st.success(f"✅ Sampled {len(sample_df)} rows, collected samples for {len(samples_map)} columns")

                    except Exception as e:
                        st.warning(f"One-pass sampling failed: {str(e)}")
                        st.info("Falling back to per-column sampling (this will be slower)...")

                        # Fallback to chunked sampling
                        columns = st.session_state.columns_data['column_name'].tolist()
                        data_types = dict(zip(
                            st.session_state.columns_data['column_name'],
                            st.session_state.columns_data['data_type']
                        ))

                        samples_map = sample_rows_chunked(
                            st.session_state.connection,
                            st.session_state.selected_db,
                            st.session_state.selected_schema,
                            st.session_state.selected_view,
                            columns,
                            data_types,
                            max_values=15,
                            max_workers=8
                        )

                        st.session_state.sample_data = samples_map
                        st.success(f"✅ Collected samples for {len(samples_map)} columns")

            except Exception as e:
                st.error(f"Sampling failed: {str(e)}")
                st.code(traceback.format_exc())

    # Show sampling status
    if skip_sampling:
        st.session_state.sample_data = {}
        st.info("ℹ️ Sampling skipped - LLM will generate descriptions without sample values")
    elif st.session_state.sample_data is not None:
        num_sampled = sum(1 for v in st.session_state.sample_data.values() if v)
        st.success(f"✅ Samples collected for {num_sampled}/{num_columns} columns")

def render_llm_settings():
    """Render LLM settings panel - Cortex only"""
    if st.session_state.columns_data is None:
        return

    st.markdown("---")
    st.header("🤖 LLM Settings")

    st.info("ℹ️ **LLM executed inside Snowflake (Cortex)** - No external API keys required")

    with st.expander("⚙️ Configure LLM Generation", expanded=True):
        # Cortex model selection (only)
        col1, col2 = st.columns(2)

        with col1:
            cortex_models = [
                'mixtral-8x7b',
                'mistral-large',
                'llama2-70b-chat',
                'mistral-7b'
            ]
            model = st.selectbox(
                "Cortex Model",
                cortex_models,
                help="Snowflake Cortex model to use for generation"
            )

        with col2:
            output_style = st.radio(
                "Output Style",
                ['technical', 'business'],
                format_func=lambda x: 'Technical/Data Engineering' if x == 'technical' else 'Business-Friendly',
                help="Choose the tone and style of generated descriptions"
            )

            max_length = st.select_slider(
                "Max Description Length",
                options=[100, 150, 200, 300, 400, 500],
                value=200,
                help="Maximum characters per column description"
            )

        # Generation scope
        st.markdown("### Generation Scope")

        scope_option = st.radio(
            "Which columns to document?",
            ['all', 'selected', 'filter'],
            format_func=lambda x: {
                'all': 'All columns',
                'selected': 'Selected columns',
                'filter': 'Columns matching filter'
            }[x],
            horizontal=True
        )

        columns_to_generate = st.session_state.columns_data['column_name'].tolist()

        if scope_option == 'selected':
            columns_to_generate = st.multiselect(
                "Select columns",
                options=columns_to_generate,
                default=columns_to_generate[:10]
            )
        elif scope_option == 'filter':
            filter_text = st.text_input(
                "Filter (text or regex)",
                help="Enter text or regex pattern to match column names"
            )
            if filter_text:
                try:
                    pattern = re.compile(filter_text, re.IGNORECASE)
                    columns_to_generate = [
                        col for col in columns_to_generate
                        if pattern.search(col)
                    ]
                    st.info(f"Matched {len(columns_to_generate)} columns")
                except re.error:
                    # Fall back to simple text matching
                    columns_to_generate = [
                        col for col in columns_to_generate
                        if filter_text.lower() in col.lower()
                    ]
                    st.info(f"Matched {len(columns_to_generate)} columns")

        # Content toggles
        st.markdown("### Include in Prompt")

        col1, col2, col3 = st.columns(3)

        with col1:
            include_types = st.checkbox("Data types", value=True)
        with col2:
            include_samples = st.checkbox(
                "Sample values",
                value=True,
                help="⚠️ May increase cost and expose data to LLM"
            )
        with col3:
            include_view_desc = st.checkbox("View description", value=True)

        # Cost controls
        st.markdown("### Cost Controls")

        col1, col2 = st.columns(2)

        with col1:
            batch_size = st.slider(
                "Batch size (columns per LLM call)",
                min_value=25,
                max_value=150,
                value=60,
                step=5,
                help="Larger batches = fewer API calls but higher tokens per call"
            )

        with col2:
            max_columns = st.number_input(
                "Max columns per run (safety cap)",
                min_value=10,
                max_value=2000,
                value=300,
                step=50,
                help="Limit total columns to process in one run"
            )

        # Dry-run token estimate
        if st.button("📊 Estimate Tokens (Dry Run)"):
            # Build a sample prompt for first batch
            sample_batch = st.session_state.columns_data.head(min(batch_size, len(columns_to_generate)))
            sample_prompt = build_llm_prompt(
                sample_batch,
                st.session_state.get('view_description', ''),
                st.session_state.sample_data or {},
                {
                    'output_style': output_style,
                    'max_length': max_length,
                    'include_types': include_types,
                    'include_samples': include_samples,
                    'include_view_desc': include_view_desc
                }
            )

            estimated_tokens = estimate_tokens(sample_prompt)
            num_batches = (len(columns_to_generate) + batch_size - 1) // batch_size
            total_tokens = estimated_tokens * num_batches

            st.info(f"""
            **Estimate (rough):**
            - ~{estimated_tokens:,} tokens per batch
            - {num_batches} batches needed
            - ~{total_tokens:,} total input tokens
            """)

        # Store settings in session state
        st.session_state.llm_settings = {
            'model': model,
            'output_style': output_style,
            'max_length': max_length,
            'batch_size': batch_size,
            'max_columns': max_columns,
            'include_types': include_types,
            'include_samples': include_samples,
            'include_view_desc': include_view_desc,
            'columns_to_generate': columns_to_generate
        }

        return columns_to_generate

def render_generation():
    """Render generation section"""
    if st.session_state.columns_data is None:
        return

    columns_to_generate = render_llm_settings()

    if not columns_to_generate:
        st.warning("No columns selected for generation")
        return

    st.markdown("---")

    # Generate button
    col1, col2 = st.columns([2, 1])

    with col1:
        generate_button = st.button(
            f"🚀 Generate Descriptions ({len(columns_to_generate)} columns)",
            type="primary",
            use_container_width=True
        )

    with col2:
        if st.session_state.failed_batches:
            retry_button = st.button(
                f"🔄 Retry Failed Batches ({len(st.session_state.failed_batches)})",
                use_container_width=True
            )
        else:
            retry_button = False

    if generate_button or retry_button:
        settings = st.session_state.llm_settings

        # Filter columns to generate
        columns_df = st.session_state.columns_data[
            st.session_state.columns_data['column_name'].isin(columns_to_generate)
        ].copy()

        # Progress tracking
        progress_bar = st.progress(0)
        status_text = st.empty()
        time_text = st.empty()

        def update_progress(batch_idx, num_batches, start_time):
            progress = (batch_idx + 1) / num_batches
            progress_bar.progress(progress)

            elapsed = time.time() - start_time
            status_text.text(f"Processing batch {batch_idx + 1}/{num_batches}")
            time_text.text(f"Elapsed: {format_elapsed_time(elapsed)}")

        try:
            # Generate descriptions using Snowflake Cortex
            results_df, failed_batches = generate_all_descriptions(
                columns_df,
                st.session_state.get('view_description', ''),
                st.session_state.sample_data or {},
                settings,
                st.session_state.connection,
                progress_callback=update_progress
            )

            st.session_state.results_df = results_df
            st.session_state.failed_batches = failed_batches
            st.session_state.generation_complete = True

            if failed_batches:
                st.warning(f"⚠️ {len(failed_batches)} batches failed. You can retry them.")
            else:
                st.success("✅ Generation complete!")

            st.rerun()

        except Exception as e:
            st.error(f"Generation failed: {str(e)}")
            st.code(traceback.format_exc())

def render_results():
    """Render results section"""
    if st.session_state.results_df is None:
        return

    st.markdown("---")
    st.header("📋 Results")

    df = st.session_state.results_df

    # Summary stats
    col1, col2, col3, col4 = st.columns(4)

    with col1:
        st.metric("Total Columns", len(df))
    with col2:
        avg_confidence = df['confidence'].mean()
        st.metric("Avg Confidence", f"{avg_confidence:.1f}/5")
    with col3:
        errors = len(df[df['description'].str.contains('<ERROR', na=False)])
        st.metric("Errors", errors)
    with col4:
        with_samples = len(df[df['sample_values'].str.len() > 0])
        st.metric("With Samples", with_samples)

    # Search/filter
    search = st.text_input("🔍 Search columns", "")

    if search:
        mask = df['column_name'].str.contains(search, case=False, na=False) | \
               df['description'].str.contains(search, case=False, na=False)
        filtered_df = df[mask]
        st.info(f"Showing {len(filtered_df)} / {len(df)} columns")
    else:
        filtered_df = df

    # Display table
    st.dataframe(
        filtered_df,
        use_container_width=True,
        height=500,
        column_config={
            'column_name': st.column_config.TextColumn('Column Name', width='medium'),
            'data_type': st.column_config.TextColumn('Data Type', width='small'),
            'description': st.column_config.TextColumn('Description', width='large'),
            'confidence': st.column_config.NumberColumn('Confidence', width='small', format='%d'),
            'sample_values': st.column_config.TextColumn('Sample Values', width='large')
        }
    )

    # Export buttons
    st.markdown("### 📥 Export")

    col1, col2 = st.columns(2)

    with col1:
        csv_data = export_to_csv(df)
        st.download_button(
            label="📄 Download CSV",
            data=csv_data,
            file_name=f"{st.session_state.selected_view}_documentation.csv",
            mime="text/csv",
            use_container_width=True
        )

    with col2:
        try:
            excel_data = export_to_excel(df)
            st.download_button(
                label="📊 Download Excel",
                data=excel_data,
                file_name=f"{st.session_state.selected_view}_documentation.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                use_container_width=True
            )
        except ImportError:
            st.warning("openpyxl not installed - Excel export unavailable")

# ============================================================================
# MAIN APP
# ============================================================================

def main():
    """Main application"""
    st.title("❄️ Snowflake Data Catalog Documentation Generator")
    st.markdown("Generate LLM-powered column documentation for Snowflake views with 200-1200 columns")

    # Render sidebar
    render_sidebar()

    # Main content - only show if connected
    if not st.session_state.connected:
        st.info("👈 Connect to Snowflake using SSO to get started")

        st.markdown("""
        ### Features
        - 🔐 **Secure SSO authentication** via external browser
        - 📊 **View selection** with database/schema/view hierarchy
        - 🎲 **Efficient sampling** for wide tables (200-1200 columns)
        - 🤖 **LLM-powered descriptions** with Snowflake Cortex (executed inside Snowflake)
        - 📦 **Batch processing** with progress tracking and error recovery
        - 📥 **Export to CSV/Excel** with formatting

        ### Getting Started
        1. Enter your Snowflake account and email in the sidebar
        2. Click "Connect (SSO)" and authenticate in your browser
        3. Select a database, schema, and view
        4. Configure LLM settings and generate documentation
        """)

        return

    # View selection
    render_view_selection()

    # Sampling
    render_sampling_section()

    # Generation
    render_generation()

    # Results
    render_results()

if __name__ == "__main__":
    main()
